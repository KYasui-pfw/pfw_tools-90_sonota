"""genpin パッケージのテスト。外（業務ポータル・rBOM・i-Reporter）にはつながない。

実行: プロジェクト直下で  python -m unittest discover tests
"""

import io
import sqlite3
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

import openpyxl
import pandas as pd

from genpin import config, db, listing, tag_excel, update


def slip(**values):
    """業務ポータルの加工伝票1行（必要なところだけ指定する）。"""
    base = {
        "denpyo_no": "H000000001", "getsuji": "2026117", "hmcd": "740-409CA30",
        "part_name": "CAM HOLDER", "required_qty": "1", "seino": "26SKJ102", "assy_no": "J102",
        "srcd": "40", "model_name": "MX-JS3B", "inch": "30.0", "gauge": "28.0",
        "customer_cd": "J0926", "customer_name": "", "skdk": "SK", "seiban_kbn": "1",
        "assy_start": "2026/11/02", "konpo_start": "2026/11/09", "due_date": "2026/11/09",
        "is_mckr_copy": False, "version": 2, "version_source": "work_plan",
        "process_str": "B1-G1-L1", "is_plating": False,
    }
    base.update(values)
    return base


class TempDbTestCase(unittest.TestCase):
    """一時フォルダの genpinhyo.db を使うテスト。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        patcher = mock.patch.multiple(
            config,
            GENPINHYO_DB=root / "genpinhyo.db",
            BACKUP_DIR=root / "backup",
        )
        patcher.start()
        self.addCleanup(patcher.stop)
        self.addCleanup(self.tmp.cleanup)
        conn = db.connect()
        db.ensure_schema(conn)
        conn.close()

    def query(self, sql, params=()):
        conn = sqlite3.connect(config.GENPINHYO_DB)
        try:
            return conn.execute(sql, params).fetchall()
        finally:
            conn.close()


class FormatTests(unittest.TestCase):
    def test_to_int_text(self):
        self.assertEqual(update.to_int_text("34.0"), "34")
        self.assertEqual(update.to_int_text(1.0), "1")
        self.assertEqual(update.to_int_text("34.5"), "34.5")
        self.assertEqual(update.to_int_text(""), "")
        self.assertEqual(update.to_int_text(None), "")
        self.assertEqual(update.to_int_text("PT"), "PT")

    def test_getsuji_gets_underscore_when_six_digits(self):
        self.assertEqual(update.to_getsuji_text("202610"), "202610_")
        self.assertEqual(update.to_getsuji_text(2026117), "2026117")
        self.assertEqual(update.to_getsuji_text("202610_"), "202610_")

    def test_slash_date(self):
        self.assertEqual(update.to_slash_date("2026-03-12T00:00:00"), "2026/03/12")
        self.assertEqual(update.to_slash_date(None), "")

    def test_make_assy_no(self):
        self.assertEqual(update.make_assy_no("26SKD002"), "D002")
        self.assertEqual(update.make_assy_no("26P09-0186"), "26P09-0186")


class FilterTests(unittest.TestCase):
    def test_stock_assy_and_numeric_change_gear_box_are_dropped(self):
        kept, excluded = update.filter_slips([
            slip(denpyo_no="H1", assy_no="A00012345"),
            slip(denpyo_no="H2", model_name="CHANGE GEAR BOX", assy_no="000001751"),
            slip(denpyo_no="H3"),
        ])
        self.assertEqual([s["denpyo_no"] for s in kept], ["H3"])
        self.assertEqual(excluded, set())   # 払出先の除外ではないので excluded_lots には入れない

    def test_srcd_rules(self):
        kept, excluded = update.filter_slips([
            slip(denpyo_no="N40", srcd="40"),
            slip(denpyo_no="N49", srcd="49"),
            slip(denpyo_no="N60", srcd="60"),
            slip(denpyo_no="S40", srcd="40", part_name="NEEDLE DIAL"),
            slip(denpyo_no="S60", srcd="60", part_name="NEEDLE DIAL"),
            slip(denpyo_no="NPT", srcd="PT"),   # 現品票では PT を捨てない
        ])
        self.assertEqual([s["denpyo_no"] for s in kept], ["N40", "S60", "NPT"])
        self.assertEqual(excluded, {"N49", "N60", "S40"})

    def test_lot_with_one_kept_row_is_not_excluded(self):
        """MCKR複製（同じロットの 40 と 60）で片方だけ残るロットは excluded_lots に入れない。"""
        kept, excluded = update.filter_slips([
            slip(denpyo_no="H9", srcd="40", part_name="WRAP DIAL"),
            slip(denpyo_no="H9", srcd="60", part_name="WRAP DIAL", is_mckr_copy=True),
        ])
        self.assertEqual(len(kept), 1)
        self.assertEqual(excluded, set())


class BuildRecordsTests(unittest.TestCase):
    def test_record_values(self):
        drawings = {"740-409CA30": {"B1": "740-409CA$$-B1", "F1": "", "L": "740-409CA30-L",
                                    "handwritten": "X@", "spare1": "S1"}}
        record = update.build_records(
            [slip(is_plating=True, assy_no="J　102")],
            {"J0926": ("ヤマヨ　ジャージー", "JAPAN")},
            drawings,
        )[0]
        self.assertEqual(record["吋"], "30")
        self.assertEqual(record["G"], "28")
        self.assertEqual(record["組立番号"], "J 102")
        self.assertEqual(record["客先名"], "ヤマヨ ジャージー")
        self.assertEqual(record["国名"], "JAPAN")
        self.assertEqual(record["工程VERSION"], "2")
        self.assertEqual(record["ﾒｯｷ"], "1.0")
        self.assertEqual(record["B1図番"], "740-409CA$$-B1")
        self.assertEqual(record["L図番"], "740-409CA30-L")
        self.assertEqual(record["手書き図番"], "X@")
        self.assertEqual(record["予備図番1"], "S1")
        self.assertEqual(record["R図番"], "")
        self.assertEqual(set(record), set(db.GENPINHYO_COLUMNS))

    def test_unknown_customer_is_blank(self):
        record = update.build_records([slip(customer_cd="ZZZ")], {}, {})[0]
        self.assertEqual((record["客先名"], record["国名"]), ("", ""))
        self.assertEqual(record["ﾒｯｷ"], "")


class RunUpdateTests(TempDbTestCase):
    def run_update(self, slips, allocations=()):
        customers = pd.DataFrame([{"客先コード": "J0926", "客先名": "客先A", "国名": "JAPAN"}])
        with mock.patch("genpin.portal_api.fetch_denpyo", return_value=slips), \
                mock.patch("genpin.portal_api.fetch_drawings", return_value=[]), \
                mock.patch("genpin.portal_api.fetch_lots", side_effect=lambda lots: [
                    {"version": 1, "process_str": "", "is_plating": False} for _ in lots]), \
                mock.patch("genpin.irepo.fetch_customers", return_value=customers), \
                mock.patch("genpin.rbom.fetch_change_gear_boxes", return_value=list(allocations)):
            return update.run_update()

    def test_writes_rows_parts_and_excluded(self):
        result = self.run_update([slip(), slip(denpyo_no="H000000002", srcd="49", part_name="NEW PART")])
        self.assertEqual(result.written, 1)
        self.assertEqual(self.query("SELECT ロット番号, 客先名 FROM genpinhyo"), [("H000000001", "客先A")])
        self.assertEqual(self.query("SELECT ロット番号 FROM excluded_lots"), [("H000000002",)])
        # 部品名は残した伝票のものだけ、QRコードレイアウト区分 2（新しい部品名の目印）で足す
        self.assertEqual(self.query("SELECT * FROM buhin_irepo_mst"), [("CAM HOLDER", 0, 2, 0)])
        self.assertEqual(self.query("SELECT * FROM delete_mst"), [("H000000001", 0)])

    def test_existing_part_settings_are_kept(self):
        conn = sqlite3.connect(config.GENPINHYO_DB)
        conn.execute("INSERT INTO buhin_irepo_mst VALUES ('CAM HOLDER', 1, 0, 1)")
        conn.commit()
        conn.close()
        result = self.run_update([slip()])
        self.assertEqual(result.new_parts, 0)
        self.assertEqual(self.query("SELECT * FROM buhin_irepo_mst"), [("CAM HOLDER", 1, 0, 1)])

    def test_lot_missing_from_denpyo_is_flagged_deleted(self):
        month = f"{config.now():%Y%m}"
        self.run_update([slip(getsuji=month + "1"), slip(denpyo_no="H000000002", getsuji=month + "1")])
        result = self.run_update([slip(getsuji=month + "1")])
        self.assertEqual((result.deleted_on, result.deleted_off), (1, 0))
        self.assertEqual(self.query("SELECT 削除flg FROM delete_mst WHERE ロット番号 = 'H000000002'"), [(1,)])
        self.assertEqual(self.query("SELECT ロット番号, old_flg, new_flg FROM flag_change_log"),
                         [("H000000002", 0, 1)])

    def test_old_months_are_not_touched(self):
        """当月〜3か月先より前のロットは、加工伝票から消えても削除扱いにしない。"""
        self.run_update([slip(getsuji="2020011")])
        self.run_update([slip(denpyo_no="H000000002")])
        self.assertEqual(self.query("SELECT 削除flg FROM delete_mst WHERE ロット番号 = 'H000000001'"), [(0,)])

    def test_change_gear_box_gets_ch_number_and_keeps_it(self):
        allocation = {"SEINO": "26SKD002", "HMCD": "G-6503GD", "ALCQTY": 1.0,
                      "TDPLNDT": "2026-11-12T00:00:00", "SPPLDT": "2026-11-20T00:00:00",
                      "ASPLDT": "2026-11-10T00:00:00", "GETSUJI": 2026117}
        self.run_update([slip(seino="26SKD002")], [allocation])
        self.run_update([slip(seino="26SKD002")], [allocation])
        rows = self.query("SELECT ロット番号, 組立番号, 機種, 必要数, 指定納期 FROM genpinhyo WHERE 部品名 = 'CHANGE GEAR BOX'")
        self.assertEqual(rows, [("CH00000001", "D002", "MX-JS3B", "1", "2026/11/12")])
        self.assertEqual(self.query("SELECT * FROM cgb_seino_ch_mapping"), [("26SKD002", "CH00000001")])

    def test_api_failure_writes_nothing(self):
        from genpin import portal_api
        with mock.patch("genpin.portal_api.fetch_denpyo", side_effect=portal_api.PortalApiError("x")):
            with self.assertRaises(portal_api.PortalApiError):
                update.run_update()
        self.assertEqual(self.query("SELECT COUNT(*) FROM genpinhyo"), [(0,)])


class ListingTests(unittest.TestCase):
    def test_build_tag_list(self):
        genpinhyo = pd.DataFrame([
            {"ロット番号": "H1", "部品名": "P", "組立番号": "J102", "工程VERSION": "1.0"},
            {"ロット番号": "H2", "部品名": "P", "組立番号": "J102", "工程VERSION": ""},
            {"ロット番号": "H3", "部品名": "P", "組立番号": "A00012345", "工程VERSION": "1"},
            {"ロット番号": "CH00000001", "部品名": "Q", "組立番号": "D002", "工程VERSION": "1"},
        ])
        parts = pd.DataFrame([{"部品名": "P", "ireporter管理flg": 1, "QRコードレイアウト区分": 0, "縦横区分": 1}])
        deletes = pd.DataFrame([{"ロット番号": "H1", "削除flg": 0}, {"ロット番号": "H2", "削除flg": 1},
                                {"ロット番号": "H3", "削除flg": 0}])
        reports = pd.DataFrame([{"帳票発行ID": 10, "ロット番号": "H1"}, {"帳票発行ID": 12, "ロット番号": "H1"}])

        df = listing.build_tag_list(genpinhyo, parts, deletes, reports)
        self.assertEqual(df["ロット番号"].tolist(), ["H1", "CH00000001"])
        self.assertEqual(df["帳票発行ID"].tolist(), [12, 0])
        self.assertEqual(df["工程VERSION"].tolist(), [1, 1])
        self.assertEqual(df["ireporter管理flg"].tolist(), [1, ""])
        self.assertEqual(listing.needs_report(df).empty, True)


class TagExcelTests(unittest.TestCase):
    def row(self, **values):
        # ⚠ 半角カナの「ﾒｯｷ」はキーワード引数に書けない（Python が全角の「メッキ」に読み替える）。
        #   ﾒｯｷ を変えるときは **{"ﾒｯｷ": ...} で渡す。
        base = {"ロット番号": "H1", "月次": "2026117", "組立番号": "J102", "機種": "M", "吋": "30", "G": "28",
                "完成部番": "740-409CA30", "部品名": "P", "客先コード": "J0926", "ﾒｯｷ": "", "必要数": "1",
                "工程": "B1-G1", "QRコードレイアウト区分": 0, "帳票発行ID": 123,
                "B1図番": "b1", "B2図番": "", "F1図番": "", "F2図番": "", "L図番": "l", "R図番": "", "C図番": ""}
        base.update(values)
        return base

    def test_qr_text(self):
        self.assertEqual(tag_excel.qr_text(self.row()), "jp.co.cimtops.ireporter.openreport:repid=123")
        self.assertEqual(tag_excel.qr_text(self.row(QRコードレイアウト区分=1, **{"ﾒｯｷ": "1.0"})),
                         "H1,2026117,J102,M,30,28,740-409CA30,P,'',J0926,1.0")

    def test_split_per_piece_and_plating(self):
        entries = tag_excel.entries_from_rows([
            self.row(完成部番="111-111A30", 必要数="3"),
            self.row(ロット番号="H2", **{"ﾒｯｷ": "1.0"}),
        ])
        self.assertEqual(len(entries), 4)
        self.assertTrue(entries[0].notice.endswith("1/3"))
        self.assertTrue(entries[2].notice.endswith("3/3"))
        self.assertEqual(entries[3].notice, "メッキ")

    def test_book_cells(self):
        data, name = tag_excel.build_tag_book(
            [self.row(), self.row(ロット番号="H2"), self.row(ロット番号="H3"), self.row(ロット番号="H4")])
        self.assertTrue(name.startswith("現品票_"))
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        rows = tag_excel.ROWS_PER_TAG
        self.assertEqual(sheet.cell(1, 30).value, "H1")
        self.assertEqual(sheet.cell(1 + rows, 30).value, "H2")
        self.assertEqual(sheet.cell(4, 24).value, "11月7次")
        self.assertEqual(sheet.cell(8, 7).value, "b1")
        self.assertEqual(len(sheet._images), 4)
        # 3枚ごとに改ページ
        self.assertEqual([b.id for b in sheet.row_breaks.brk], [rows * 3])
        # ページの中央に寄せる。最後のページ（1枚）も3枚分の高さを印刷範囲にして、位置をそろえる
        self.assertTrue(sheet.print_options.horizontalCentered)
        self.assertTrue(sheet.print_options.verticalCentered)
        self.assertIn(f"$AR${rows * 6}", sheet.print_area)
        self.assertEqual(sheet.row_dimensions[rows * 6].height, sheet.row_dimensions[rows].height)

    def test_three_tags_fit_a4(self):
        """1ページ分（3枚 − 最後のすき間）の高さが A4 に収まり、すき間が点線で2等分されていること。"""
        from genpin import tag_layout
        heights = tag_layout.ROW_HEIGHTS
        self.assertEqual(tag_layout.TAGS_PER_PAGE, 3)
        self.assertEqual(len(heights), tag_layout.ROWS_PER_TAG)
        gap = heights[tag_layout.CUT_ROW - 1:]
        self.assertEqual(gap[0], gap[1])
        # RICOH IM C3500 で 891.5pt ではみ出した（875.5pt までは収まった）。それより余裕を残しておく。
        self.assertLessEqual(sum(heights) * 3 - sum(gap), 865)

    def test_last_tag_of_page_hides_gap(self):
        """ページの最後の1枚だけ、下のすき間（19・20行目）を隠す。中央寄せで切った3枚を等分にするため。"""
        data, _ = tag_excel.build_tag_book([self.row(ロット番号=f"H{i}") for i in range(4)])
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        rows = tag_excel.ROWS_PER_TAG
        hidden = [row for row in range(1, rows * 6 + 1) if sheet.row_dimensions[row].hidden]
        self.assertEqual(hidden, [rows * 2 + 19, rows * 2 + 20, rows * 5 + 19, rows * 5 + 20])

    def test_borders_survive_merge(self):
        """結合したセルの罫線が、2枚目にも残っていること（結合で罫線が消えていた不具合）。"""
        data, _ = tag_excel.build_tag_book([self.row(), self.row(ロット番号="H2")])
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        rows = tag_excel.ROWS_PER_TAG
        for offset in (0, rows):
            self.assertEqual(sheet.cell(1 + offset, 5).border.top.style, "thin")       # 現品票の上辺
            self.assertEqual(sheet.cell(2 + offset, 40).border.bottom.style, "thin")   # ロット番号の下辺
            self.assertEqual(sheet.cell(18 + offset, 30).border.bottom.style, "thin")  # 工程の下辺
            self.assertEqual(sheet.cell(17 + offset, 44).border.right.style, "thin")   # 工程の右辺

    def test_layout_labels(self):
        """図番の見出し（B・F の番号は無し、手書図番・予備項目を追加）と、2枚目にも枠があること。"""
        data, _ = tag_excel.build_tag_book([self.row(**{"手書き図番": "X@", "予備図番1": "S1"}), self.row(ロット番号="H2")])
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        labels = [sheet.cell(row, 1).value for row in range(8, 17)]
        self.assertEqual(labels, ["B　図番：", "B　図番：", "F　図番：", "F　図番：",
                                  "L　図番：", "R　図番：", "C　図番：", "手書図番：", "予備項目："])
        self.assertEqual(sheet.cell(5, 1).value, "完成部番：")
        self.assertEqual((sheet.cell(15, 7).value, sheet.cell(16, 7).value), ("X@", "S1"))
        rows = tag_excel.ROWS_PER_TAG
        self.assertEqual(sheet.cell(16 + rows, 1).value, "予備項目：")
        self.assertEqual(sheet.cell(5 + rows, 7).font.sz, sheet.cell(5, 7).font.sz)
        self.assertEqual(sheet.cell(19, 1).border.bottom.style, "dotted")   # 切り取り線

    def test_qr_can_be_read(self):
        import zxingcpp
        from PIL import Image
        text = "jp.co.cimtops.ireporter.openreport:repid=68376"
        image = Image.open(io.BytesIO(tag_excel.make_qr_png(text)))
        self.assertEqual(image.size, (135, 135))
        background = Image.new("RGB", image.size, "white")
        background.paste(image, mask=image.split()[3])
        found = zxingcpp.read_barcodes(background.resize((540, 540), Image.NEAREST))
        self.assertEqual(found[0].text, text)

    def test_single_tag(self):
        entry = tag_excel.TagEntry(lot_no="KARI1", getsuji="11月_次", hmcd="X", qty=1, header="【再発行】",
                                   qr_png=tag_excel.make_qr_png("x"))
        data, name = tag_excel.build_single_tag(entry)
        sheet = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertTrue(name.startswith("手入力現品票_"))
        self.assertEqual((sheet.cell(1, 30).value, sheet.cell(1, 15).value), ("KARI1", "【再発行】"))


class RebuildTests(unittest.TestCase):
    """壊れた genpinhyo の作り直し（db.rebuild_genpinhyo）。"""

    def test_keeps_latest_row_per_lot(self):
        conn = sqlite3.connect(":memory:")
        # 索引が壊れて同じロットが2行入った状態を、主キーの無いテーブルで再現する。
        conn.execute(f"CREATE TABLE genpinhyo ({', '.join(c + ' TEXT' for c in db.GENPINHYO_COLUMNS)})")
        conn.executemany("INSERT INTO genpinhyo (ロット番号, B1図番) VALUES (?, ?)",
                         [("H1", "古い"), ("H2", "x"), ("H1", "新しい")])
        conn.commit()

        scanned, after = db.rebuild_genpinhyo(conn)
        self.assertEqual((scanned, after), (3, 2))
        self.assertEqual(conn.execute("SELECT ロット番号, B1図番 FROM genpinhyo ORDER BY 1").fetchall(),
                         [("H1", "新しい"), ("H2", "x")])
        self.assertEqual(conn.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        # 作り直したあとは主キーが効く（同じロットは REPLACE で1行のまま）
        conn.execute("REPLACE INTO genpinhyo (ロット番号) VALUES ('H1')")
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM genpinhyo WHERE ロット番号 = 'H1'").fetchone(), (1,))


class CorruptedDbTests(TempDbTestCase):
    run_update = RunUpdateTests.run_update

    def test_corrupted_db_stops_before_writing(self):
        with mock.patch("genpin.db.check_integrity", side_effect=db.DatabaseCorruptedError("壊れている")):
            with self.assertRaises(db.DatabaseCorruptedError):
                self.run_update([slip()])
        self.assertEqual(self.query("SELECT COUNT(*) FROM genpinhyo"), [(0,)])
        self.assertFalse(config.BACKUP_DIR.exists() and any(config.BACKUP_DIR.iterdir()))

    def test_healthy_db_passes_check(self):
        conn = db.connect()
        try:
            db.check_integrity(conn)
        finally:
            conn.close()


class BackupTests(TempDbTestCase):
    def test_backup_and_prune(self):
        config.BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        old = config.BACKUP_DIR / "20200101_genpinhyo.db"
        old.write_bytes(b"")
        db.backup_genpinhyo()
        names = sorted(p.name for p in config.BACKUP_DIR.iterdir())
        self.assertEqual(names, [f"{config.now():%Y%m%d}_genpinhyo.db"])


if __name__ == "__main__":
    unittest.main()
