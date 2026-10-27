"""現品票の Excel（QRコード付き）を作る。

    一覧から複数枚   build_tag_book      （現品票発行の「チェック済の印刷」）
    1枚だけ          build_single_tag    （手入力発行・チェックシート再作成）

⚠ 2026-10-27 から、ひな形の Excel は使わない。1枚分の枠（罫線・見出し）は tag_layout.py が
  コードで描き、ここは値と QR を入れるだけ。見た目を変えるときは tag_layout.py を直す。

⚠ QR の画像は以前は work/ フォルダに書き出してから Excel に貼っていた（20分ごとに掃除していた）。
  いまはメモリの中だけで作るので、ファイルは残らない。
"""

import io
from dataclasses import dataclass

import openpyxl
import qrcode
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image, ImageOps

from . import config, tag_layout

# 1枚の現品票が何行ぶんか（2枚目は20行下から始まる）。
ROWS_PER_TAG = tag_layout.ROWS_PER_TAG

# 現品票の各項目を書くセル（1枚目の行, 列）。2枚目以降は行に ROWS_PER_TAG × n を足す。
TAG_CELLS = {
    "header": (1, 15),       # 【再発行】などの見出し（再発行のときだけ）
    "notice": (2, 10),       # 注意書き（メッキ・1/3 など）
    "issued": (3, 1),        # 発行日
    "qty": (4, 21),          # 個数
    "hmcd": (5, 7),          # 完成部番
    "B1": (8, 7),            # 図番（見出しは tag_layout.DRAWING_ROWS）
    "B2": (9, 7),
    "F1": (10, 7),
    "F2": (11, 7),
    "L": (12, 7),
    "R": (13, 7),
    "C": (14, 7),
    "handwritten": (15, 7),  # 手書図番（末尾が @ の図番）
    "spare1": (16, 7),       # 予備項目（7つの枠からあふれた図番の1つ目）
    "lot_no": (1, 30),       # ロット番号
    "getsuji": (4, 24),      # 月次（例 '11月7次'）
    "assy_no": (7, 24),      # 組立番号
    "model": (10, 24),       # 機種
    "inch": (13, 32),
    "gauge": (13, 38),
    "process": (tag_layout.PROCESS_ROW, 6),  # 工程
}
QR_ANCHOR_COLUMN = "AI"      # QR を貼る列（行は 3 行目）
QR_ANCHOR_ROW = 3
QR_VERSION = 5               # QR の型番（大きさ）。昔から5
QR_BOX_SIZE = 3              # 1マスの画素数
QR_BORDER = 4                # まわりの余白（マス）

# 現品票に載せる図番（TAG_CELLS のキー → 一覧の列名）。
DRAWING_FIELDS = {
    "B1": "B1図番", "B2": "B2図番", "F1": "F1図番", "F2": "F2図番",
    "L": "L図番", "R": "R図番", "C": "C図番",
    "handwritten": "手書き図番", "spare1": "予備図番1",
}

# 完成部番にこれを含み、必要数が2以上なら、1個ずつ別の現品票にする（1/3, 2/3, 3/3）。
SPLIT_PER_PIECE = ("-111", "-404", "-102")

# QRコードレイアウト区分。
QR_LAYOUT_IREPORTER = 0   # i-Reporter の帳票を開く QR
QR_LAYOUT_LEGACY = 1      # 昔の形（項目をカンマでつないだ文字）
QR_LAYOUT_NEW_PART = 2    # 新しい部品名でまだ設定していない（QR の代わりにエラーの画像）

PLATING_NOTICE = "メッキ"


@dataclass
class TagEntry:
    """現品票1枚ぶんの中身。"""
    lot_no: str
    getsuji: str          # 印字する月次（例 '11月7次'）
    hmcd: str
    qty: str
    assy_no: str = ""
    model: str = ""
    inch: str = ""
    gauge: str = ""
    process: str = ""
    notice: str = ""
    header: str = ""
    drawings: dict = None   # {"B1": ..., ..., "C": ..., "handwritten": ..., "spare1": ...}
    qr_png: bytes = b""     # QR（またはエラー画像）の PNG


def is_plating(value):
    """genpinhyo.ﾒｯｷ の値がメッキありか（'1.0' や 1）。"""
    try:
        return value not in ("", None) and int(float(value)) == 1
    except (TypeError, ValueError):
        return False


def getsuji_label(getsuji):
    """月次 '2026117' → '11月7次'。"""
    getsuji = str(getsuji)
    return f"{getsuji[4:6]}月{getsuji[6:]}次"


def qr_text(row):
    """一覧の1行から、QR に入れる文字を作る。"""
    if row.get("QRコードレイアウト区分") == QR_LAYOUT_LEGACY:
        return (
            f"{row['ロット番号']},{row['月次']},{row['組立番号']},{row['機種']},{row['吋']},"
            f"{row['G']},{row['完成部番']},{row['部品名']},'',{row['客先コード']},{row['ﾒｯｷ']}"
        )
    return ireporter_qr_text(row.get("帳票発行ID", 0))


def ireporter_qr_text(report_id):
    """i-Reporter の帳票を開く QR の文字。"""
    return f"jp.co.cimtops.ireporter.openreport:repid={report_id}"


def make_qr_png(text):
    """QR の PNG を作る（白い部分は透明にする。現品票の枠線に重ねても見えるように）。

    大きさは昔と同じ（型番5・誤り訂正L・1マス3画素・余白4マス＝135×135画素）。
    ⚠ 2026-10-27 に pyqrcode から qrcode に替えた（約2倍速い）。白黒の模様は少し変わることが
      あるが、読み取れる中身は同じ（tests で読み取って確かめている）。
    """
    code = qrcode.QRCode(version=QR_VERSION, error_correction=qrcode.constants.ERROR_CORRECT_L,
                         box_size=QR_BOX_SIZE, border=QR_BORDER)
    code.add_data(text.encode("utf-8"))
    code.make(fit=False)
    gray = code.make_image().get_image().convert("L")

    # 黒いところだけ見えるようにする（黒＝不透明、白＝透明）。
    image = Image.new("RGBA", gray.size, (0, 0, 0, 255))
    image.putalpha(ImageOps.invert(gray))

    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def error_png():
    """新しい部品名のときに QR の代わりに貼る画像。"""
    return (config.TEMPLATE_DIR / "err.png").read_bytes()


def entries_from_rows(rows):
    """一覧の行（辞書）から、現品票の中身のリストを作る。

    ・QRコードレイアウト区分が 2（新しい部品名）なら、QR の代わりにエラーの画像
    ・完成部番が -111/-404/-102 を含み、必要数が2以上なら、1個ずつ別の現品票（1/3 など）
    ・メッキありなら注意書きに「メッキ」（1個ずつ分けるときは 1/3 などが優先）
    """
    entries = []
    for row in rows:
        drawings = drawings_of(row)
        common = dict(
            lot_no=row["ロット番号"],
            getsuji=getsuji_label(row["月次"]),
            hmcd=row["完成部番"],
            qty=row["必要数"],
            assy_no=row["組立番号"],
            model=row["機種"],
            inch=row["吋"],
            gauge=row["G"],
            process=row["工程"],
            drawings=drawings,
        )
        plating_notice = PLATING_NOTICE if is_plating(row.get("ﾒｯｷ")) else ""

        if row.get("QRコードレイアウト区分") == QR_LAYOUT_NEW_PART:
            entries.append(TagEntry(notice=plating_notice, qr_png=error_png(), **common))
            continue

        png = make_qr_png(qr_text(row))
        count = _piece_count(row)
        if any(code in str(row["完成部番"]) for code in SPLIT_PER_PIECE) and count > 1:
            for number in range(1, count + 1):
                # 先頭の全角スペースは、注意書きのセルの右寄りに表示させるため（昔からの見た目）。
                notice = "　" * 12 + f"{number}/{count}"
                entries.append(TagEntry(notice=notice, qr_png=png, **common))
        else:
            entries.append(TagEntry(notice=plating_notice, qr_png=png, **common))
    return entries


def drawings_of(row):
    """一覧の行（辞書）から、現品票に載せる図番を取り出す。"""
    return {key: row.get(column) or "" for key, column in DRAWING_FIELDS.items()}


def _piece_count(row):
    try:
        return int(float(row["必要数"]))
    except (TypeError, ValueError):
        return 0


def _write_entry(sheet, entry, issued, offset):
    """1枚ぶんを書く。offset は行のずらし幅（n 枚目なら ROWS_PER_TAG × n）。"""
    values = {
        "notice": entry.notice,
        "issued": issued,
        "qty": entry.qty,
        "hmcd": entry.hmcd,
        "lot_no": entry.lot_no,
        "getsuji": entry.getsuji,
        "assy_no": entry.assy_no,
        "model": entry.model,
        "inch": entry.inch,
        "gauge": entry.gauge,
        "process": entry.process,
    }
    values.update(entry.drawings or {})
    if entry.header:
        values["header"] = entry.header

    for key, value in values.items():
        row, column = TAG_CELLS[key]
        sheet.cell(row + offset, column).value = value

    image = ExcelImage(io.BytesIO(entry.qr_png))
    sheet.add_image(image, f"{QR_ANCHOR_COLUMN}{QR_ANCHOR_ROW + offset}")


def _save(book):
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_tag_book(rows, shifted=False):
    """一覧の行から、現品票の Excel を作る（必要な枚数ぶんだけ）。

    Args:
        rows: 一覧の行（辞書）のリスト。並び順のまま印字する
        shifted: True なら「ずれる場合」の印刷設定（少し縮めて上の余白を詰める）

    Returns:
        (Excel のバイト列, ダウンロードのファイル名)
    """
    now = config.now()
    issued = f"{now:%Y/%m/%d}"
    entries = entries_from_rows(rows)

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "現品票"
    tag_layout.setup_sheet(sheet, len(entries), shifted=shifted)

    tag_layout.draw_tags(sheet, len(entries))
    if not entries:
        sheet.cell(*TAG_CELLS["lot_no"]).value = "対象データなし"

    for index, entry in enumerate(entries):
        _write_entry(sheet, entry, issued, ROWS_PER_TAG * index)

    prefix = "現品票B_" if shifted else "現品票_"
    return _save(book), f"{prefix}{now:%Y%m%d%H%M%S}.xlsx"


def build_single_tag(entry):
    """現品票を1枚だけ作る（手入力発行・チェックシート再作成）。

    Returns:
        (Excel のバイト列, ダウンロードのファイル名)
    """
    now = config.now()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "手入力現品票"
    tag_layout.setup_sheet(sheet, 1)
    tag_layout.draw_tags(sheet, 1)
    _write_entry(sheet, entry, f"{now:%Y/%m/%d}", 0)
    return _save(book), f"手入力現品票_{now:%Y%m%d%H%M%S}.xlsx"
