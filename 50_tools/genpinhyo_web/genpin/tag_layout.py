"""現品票1枚分のレイアウト（罫線・結合・文字の大きさ・見出し）。Excel のひな形を使わずにコードで作る。

⚠ 2026-10-27 まではひな形の Excel（template*.xlsx が16ファイル。最大4,320枚分・12.5MB）に
  値を入れていた。表記を変える・行を足すには16ファイルすべてを直す必要があったので、
  レイアウトをここに書き、必要な枚数だけ並べる形に変えた（利用者の依頼）。
  **見た目を変えるときは、このファイルだけを直す。**

【1枚の形】45列（列幅 2.25 の方眼）× 20行。A4 縦。

    行  左（A〜W）                           右（X〜AR）
    1-2 現 品 票 ｜注意事項：（中身は 2-3 行）  ロット番号：｜（ロット番号）
    3   （発行日）                           （QR を AI3 に貼る）
    4   ｜個数：（個数）                      （月次 4-6）
    5-6 完成部番：（完成部番）
    7                                       （組立番号 7-9）
    8-9   B　図番：（図番）
    10-11 F　図番：（図番）                   （機種 10-12）
    12  L　図番：
    13  R　図番：                            （吋）" /（ゲージ）G  13-14
    14  C　図番：
    15  手書図番：（末尾が @ の図番）
    16  予備項目：（予備1）
    17-18 工程：（工程）
    19  （すき間の上半分。下辺が切り取り線＝点線）
    20  （すき間の下半分）  ※ページの最後の1枚は 19-20 を隠す

⚠ 値を入れるセルの位置は tag_excel.TAG_CELLS が持っている。行を足したらそちらも直す。
"""

from copy import copy

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.merge import MergedCellRange
from openpyxl.worksheet.pagebreak import Break

FONT_NAME = "ＭＳ Ｐゴシック"
COLUMN_COUNT = 45            # A〜AS
COLUMN_WIDTH = 2.25          # 方眼の1マスの幅
ROWS_PER_TAG = 20

# 行の高さ（pt）。上から1行目〜20行目。
# ⚠ 点線で切ると A4 がちょうど3等分になり、枠がそれぞれの真ん中に来るようにしてある（利用者の希望 2026-10-27）。
#   ・19行目と20行目（すき間）は同じ高さで、その境目に点線を引く
#   ・ページの最後の1枚はすき間を隠し（setup_sheet）、残りをページの上下中央に寄せる
#   → ページの上下の余白も「すき間の半分」になるので、切った3枚が同じ大きさになる。
#   1枚 303pt は、印刷で約0.926倍（RICOH IM C3500。Excel が画面の解像度で行の高さを丸めるため）されて
#   約 280.5pt ＝ A4 の高さ 841.7pt の3分の1。実測で 280.0 / 280.5 / 281.2pt に切れた。
#   図番は9行（13.5pt＝11pt の文字の標準）。昔は図番7行（17.25pt）・すき間 27＋27.75pt だった。
#   ⚠ 高さを増やすと Excel が途中で勝手に改ページし、現品票が上下に切れて以降が全部ずれる。
#     1ページ分（3枚 − 最後のすき間）は今 860.5pt。同じプリンターで 891.5pt にすると切れた（875.5pt までは収まった）。
ROW_HEIGHTS = [15, 15, 15, 18, 15, 15, 10] + [13.5] * 9 + [15, 15, 24.25, 24.25]

# 図番の行（見出し, 行）。上から順に並ぶ。見出しの全角スペースは昔の見た目に合わせている。
DRAWING_ROWS = [
    ("B　図番：", 8),
    ("B　図番：", 9),
    ("F　図番：", 10),
    ("F　図番：", 11),
    ("L　図番：", 12),
    ("R　図番：", 13),
    ("C　図番：", 14),
    ("手書図番：", 15),
    ("予備項目：", 16),
]
PROCESS_ROW = 17             # 工程（2行ぶん）
CUT_ROW = 19                 # 切り取り線（点線）
RIGHT_EDGE_LAST_ROW = 16     # 右端（AR）の縦線を引く最後の行

# 印刷の設定（インチ）。ページの水平・垂直の中央に寄せる（2026-10-27 利用者と決定）。
# 余白は「これ以上は詰めない」最小の幅で、実際の位置は中央寄せで決まる。
# ⚠ 上下は 0 にしておく。1cm などにすると、Excel が「余白の内側に入りきらない」として上に寄せてしまい、
#   上下の余白（＝切った1枚目・3枚目の大きさ）がそろわない（上 1cm で 上 28pt・下 16pt になった）。
# 「ずれる場合」は全体を少し縮める（昔のひな形 *b.xlsx と同じ 97%）。縮めたぶん、切った3枚は少し不ぞろいになる。
MARGIN_LEFT = MARGIN_RIGHT = 0.2362204724409449
MARGIN_TOP = 0
MARGIN_BOTTOM = 0
SHIFTED_SCALE = 97
PAPER_A4 = 9

# 1ページに何枚ずつ印刷するか（その枚数ごとに改ページを入れる）。昔のひな形と同じ3枚。
# ⚠ ROW_HEIGHTS の合計×3 が A4 の印刷できる高さに収まっていること。収まらないと Excel が
#   途中で勝手に改ページし、現品票が上下に切れる（昔のひな形は改ページを入れておらず、そうなることがあった）。
TAGS_PER_PAGE = 3

THIN = Side(style="thin")
DOTTED = Side(style="dotted")


def _font(size, bold=False):
    return Font(name=FONT_NAME, size=size, bold=bold)


def _cells(ws, area, offset):
    """'A1:I2' のような範囲を、offset 行ずらしたセルの並び（行ごと）にする。"""
    start, _, end = area.partition(":")
    end = end or start
    first_col, first_row = _split(start)
    last_col, last_row = _split(end)
    return [
        [ws.cell(row + offset, col) for col in range(first_col, last_col + 1)]
        for row in range(first_row, last_row + 1)
    ]


def _split(ref):
    letters = "".join(ch for ch in ref if ch.isalpha())
    return column_index_from_string(letters), int(ref[len(letters):])


def _merge(ws, min_row, min_col, max_row, max_col):
    """セルを結合する。

    ⚠ openpyxl の ws.merge_cells は、結合を1つ足すたびに「既にある結合すべて」と重なりを
      照らし合わせる。現品票は1枚に結合が約35か所あり、1,800枚だと6万か所になって
      照らし合わせが何十億回にもなる（10分たっても終わらなかった）。
      現品票どうしの結合は重ならないことが分かっているので、照らし合わせを省いて足す。
    """
    merged = MergedCellRange(ws, CellRange(min_col=min_col, min_row=min_row,
                                           max_col=max_col, max_row=max_row).coord)
    ws.merged_cells.ranges.add(merged)
    ws._clean_merge_range(merged)


def _add_border(cell, **sides):
    """セルの罫線に、指定した辺だけを足す（もともとある辺は残す）。"""
    current = cell.border
    cell.border = Border(
        left=sides.get("left", current.left),
        right=sides.get("right", current.right),
        top=sides.get("top", current.top),
        bottom=sides.get("bottom", current.bottom),
    )


def box(ws, area, offset, value=None, size=11, bold=False, horizontal=None, vertical="center",
        top=False, bottom=False, left=False, right=False, merge=True):
    """範囲を結合して、文字の書式と外周の罫線を付ける。

    文字が枠より長いときは縮めて入れる（shrink_to_fit。昔のひな形と同じ）。
    """
    rows = _cells(ws, area, offset)
    first = rows[0][0]
    if merge and (len(rows) > 1 or len(rows[0]) > 1):
        _merge(ws, first.row, first.column, rows[-1][-1].row, rows[-1][-1].column)
        # ⚠ 結合すると左上以外のセルが作り直される。結合前に取ったセルに罫線を付けても
        #   シートには残らない（罫線が途切れて出ていた）ので、取り直す。
        rows = _cells(ws, area, offset)
    if value is not None:
        first.value = value
    for row in rows:
        for cell in row:
            cell.font = _font(size, bold)
            cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, shrink_to_fit=True)
    if top:
        for cell in rows[0]:
            _add_border(cell, top=THIN)
    if bottom:
        for cell in rows[-1]:
            _add_border(cell, bottom=THIN)
    if left:
        for row in rows:
            _add_border(row[0], left=THIN)
    if right:
        for row in rows:
            _add_border(row[-1], right=THIN)


def draw_tag(ws, offset):
    """1枚分の枠と見出しを書く。offset は行のずらし幅（n 枚目なら ROWS_PER_TAG × n）。

    行の高さは setup_sheet が付ける（ページの最後の1枚だけ、下のすき間を隠すため）。
    """
    # --- 1〜3行目: 見出し・注意事項・ロット番号・発行日 ---
    box(ws, "A1:I2", offset, "現 品 票", 22, True, "center", top=True, left=True)
    box(ws, "J1:N1", offset, "注意事項：", 11, horizontal="center", vertical="top", top=True, left=True)
    box(ws, "O1:W1", offset, None, 11, horizontal="center", top=True, right=True)   # 【再発行】など
    box(ws, "J2:W3", offset, None, 14, horizontal="center", bottom=True, left=True, right=True)
    box(ws, "X1:AC2", offset, "ロット番号：", 12, horizontal="center", top=True, bottom=True, left=True)
    box(ws, "AD1:AR2", offset, None, 16, True, "left", top=True, bottom=True, right=True)
    box(ws, "A3:I3", offset, None, 12, True, "center", left=True)

    # --- 4〜6行目: 個数・月次・完成部番 ---
    box(ws, "A4:O4", offset, None, top=True, left=True, merge=False)
    box(ws, "P4:T4", offset, "個数：", 11, horizontal="center", top=True)
    box(ws, "U4:W4", offset, None, 14, True, "left", top=True, right=True)
    box(ws, "X4:AG6", offset, None, 36, True, "left", left=True)
    box(ws, "A5:F6", offset, "完成部番：", 12, horizontal="center", left=True)
    box(ws, "G5:W6", offset, None, 22, True, "left", right=True)

    # --- 7行目: すき間・組立番号 ---
    box(ws, "A7", offset, left=True, merge=False)
    box(ws, "W7", offset, right=True, merge=False)
    box(ws, "X7:AG9", offset, None, 26, True, "left", vertical="top", left=True)

    # --- 8〜16行目: 図番 ---
    last_row = DRAWING_ROWS[-1][1]
    for label, row in DRAWING_ROWS:
        is_last = row == last_row
        # 図番の行は 13.5pt なので文字は 11pt（昔は 17.25pt の行に 12pt）。
        box(ws, f"A{row}:F{row}", offset, label, 11, horizontal="center", left=True, bottom=is_last)
        box(ws, f"G{row}:W{row}", offset, None, 11, True, "left", right=True, bottom=is_last)

    # --- 右側: 機種・吋/ゲージ ---
    # 行を詰めたぶん文字も小さくしてある（昔は機種 36pt・吋/ゲージ 24pt）。
    box(ws, "X10:AR12", offset, None, 32, True, "left", left=True, right=True)
    box(ws, "AD13:AE14", offset, None, 20, True, bottom=True, merge=False)
    box(ws, "AF13:AH14", offset, None, 20, True, "center", bottom=True)        # 吋
    box(ws, "AI13", offset, '"', 14, True, merge=False)
    box(ws, "AI14", offset, None, 12, True, bottom=True, merge=False)
    box(ws, "AJ13:AK14", offset, "/", 18, True, "center", bottom=True)
    box(ws, "AL13:AN14", offset, None, 20, True, "center", bottom=True)        # ゲージ
    box(ws, "AO13:AQ14", offset, "G", 18, True, "left", bottom=True)
    box(ws, f"AR3:AR{RIGHT_EDGE_LAST_ROW}", offset, right=True, merge=False)

    # --- 工程・切り取り線 ---
    box(ws, f"A{PROCESS_ROW}:E{PROCESS_ROW + 1}", offset, "工程：", 14, horizontal="center",
        top=True, bottom=True, left=True)
    box(ws, f"F{PROCESS_ROW}:AR{PROCESS_ROW + 1}", offset, None, 18, horizontal="left",
        top=True, bottom=True, right=True)
    for cell in _cells(ws, f"A{CUT_ROW}:AR{CUT_ROW}", offset)[0]:
        _add_border(cell, bottom=DOTTED)


def draw_tags(ws, count):
    """count 枚分の枠を描く。1枚目だけを draw_tag で描き、2枚目以降は1枚目を写す。

    ⚠ openpyxl は書式（文字・罫線）を設定するたびに、既にある書式と照らし合わせる。
      1枚ずつ draw_tag で描くと1枚あたり約0.16秒かかり、1,800枚で5分近くになる。
      1枚目のセルの書式（_style）をそのまま写せば、照らし合わせが要らないので速い。
    """
    count = max(count, 1)
    draw_tag(ws, 0)
    template = [
        [ws.cell(row, col) for col in range(1, COLUMN_COUNT + 1)]
        for row in range(1, ROWS_PER_TAG + 1)
    ]
    first_merges = [
        (r.min_row, r.min_col, r.max_row, r.max_col)
        for r in ws.merged_cells.ranges if r.max_row <= ROWS_PER_TAG
    ]
    for index in range(1, count):
        offset = ROWS_PER_TAG * index
        # ⚠ 結合を先にする。結合すると、左上以外のセルが作り直されて書式が消えるため。
        for min_row, min_col, max_row, max_col in first_merges:
            _merge(ws, min_row + offset, min_col, max_row + offset, max_col)
        for row_cells in template:
            for source in row_cells:
                target = ws.cell(source.row + offset, source.column)
                target._style = copy(source._style)
                if source.value is not None and not isinstance(source, MergedCell):
                    target.value = source.value


def setup_sheet(ws, tag_count, shifted=False):
    """シートの共通設定（方眼・印刷の向きと余白・印刷範囲・改ページ）。"""
    ws.sheet_view.showGridLines = False
    ws.sheet_format.defaultColWidth = COLUMN_WIDTH
    ws.sheet_format.defaultRowHeight = 13.5
    for col in range(1, COLUMN_COUNT + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = COLUMN_WIDTH

    ws.page_setup.paperSize = PAPER_A4
    ws.page_setup.orientation = "portrait"
    ws.page_margins.left = MARGIN_LEFT
    ws.page_margins.right = MARGIN_RIGHT
    ws.page_margins.top = MARGIN_TOP
    ws.page_margins.bottom = MARGIN_BOTTOM
    ws.page_margins.header = 0
    ws.page_margins.footer = 0
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True
    if shifted:
        ws.page_setup.scale = SHIFTED_SCALE

    count = max(tag_count, 1)
    for page_end in range(TAGS_PER_PAGE, count, TAGS_PER_PAGE):
        ws.row_breaks.append(Break(id=ROWS_PER_TAG * page_end))

    # 最後のページの足りない枚数ぶんも、同じ高さの空の行を印刷範囲に入れる。
    # ⚠ 垂直の中央寄せは「印刷範囲の高さ」で位置を決める。1〜2枚だけのページをそのまま中央に寄せると、
    #   現品票がページの真ん中あたりに出て、3枚のページと位置がそろわない。空の行は何も印刷されない。
    pages = -(-count // TAGS_PER_PAGE)
    slots = TAGS_PER_PAGE * pages
    for slot in range(slots):
        offset = ROWS_PER_TAG * slot
        page_last = slot % TAGS_PER_PAGE == TAGS_PER_PAGE - 1
        for index, height in enumerate(ROW_HEIGHTS, start=1):
            dimension = ws.row_dimensions[index + offset]
            dimension.height = height
            # ページの最後の1枚は、下のすき間（切り取り線とその下）を隠す。
            # 残りを中央に寄せると、ページの上下の余白が「すき間の半分」になり、点線で切った3枚が同じ大きさになる。
            if page_last and index >= CUT_ROW:
                dimension.hidden = True
    ws.print_area = f"A1:AR{ROWS_PER_TAG * slots}"
