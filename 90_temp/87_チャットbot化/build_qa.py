"""課題管理表Excelの「問合せ」「課題」シートを、1件1ファイルのMarkdownに変換し、索引を作る。

使い方:
    python -I build_qa.py [xlsmパス] [出力フォルダ]

依存: openpyxl
出力: qa/ 配下（INDEX.md、シートごとのフォルダと _INDEX.md、1件ごとの .md）
対象: 状況が「完了」の行のみ（回答済み）。仕掛中・未着手・空欄は出力しない。
匿名化: 発信者・対処者・確認者の列は出力しない。本文中の同じ氏名は「○○」に置き換える。
        シート名と番号（問合せ=QA-n／課題=No.n）は残し、どのシートの何番か特定できるようにする。
"""
import re
import shutil
import sys
import unicodedata
import warnings
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
XLSM = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "福原つながるプロジェクト_課題管理表.xlsm"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "qa"

TARGET_STATUS = "完了"  # 回答済みとみなす状況
SHEETS = {  # シート名 -> (ID接頭辞, 補足シートの接頭辞)
    "問合せ": ("QA-", "QA"),
    "課題": ("No.", "NO"),
}
# 列番号（0始まり）
C_NO, C_DATE, C_DEPT, C_FROM, C_KIND, C_PRIO, C_FUNC, C_TITLE, C_BODY = range(9)
C_DUE, C_OWNER, C_ANSWER, C_RESULT, C_START, C_RELEASE, C_DONE, C_CHECKER, C_MARK, C_STATUS = range(9, 19)

# 氏名の列から拾った語のうち、人名ではないもの
NOT_NAME = {"打合せ", "打ち合わせ", "うちあ合わせ", "メンテ", "福原", "購買", "購買課", "DX推進室"}
ORG_PREFIX = re.compile(r"^(PFW|DXT|DAIKO|hic|IDES|DX推進室|購買課)")
BAD_FILENAME = str.maketrans({c: "_" for c in '\\:*?"<>|/'})


def normalize(s):
    """全角英数字と半角カナだけ正規化する（grepで当たるように）。①などの記号はそのまま残す。"""
    out = []
    for ch in s:
        o = ord(ch)
        if 0xFF10 <= o <= 0xFF19 or 0xFF21 <= o <= 0xFF3A or 0xFF41 <= o <= 0xFF5A or 0xFF61 <= o <= 0xFF9F:
            ch = unicodedata.normalize("NFKC", ch)
        out.append(ch)
    s = "".join(out)
    # 半角カナ＋濁点は2文字に分かれているので、合成し直す
    return unicodedata.normalize("NFC", s)


def text(v):
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return str(v).strip()


def is_blank(v):
    return text(v) in ("", "-", "ー", "－")


def collect_names(rows):
    """発信者・対処者・確認者の列から人名を集める。"""
    names = set()
    for r in rows:
        for i in (C_FROM, C_OWNER, C_CHECKER):
            for t in re.split(r"[\s\u3000/・,、，＆&（）()※]+", text(r[i])):
                t = ORG_PREFIX.sub("", t)
                t = re.sub(r"(様|さま|さん)$", "", t)
                if re.fullmatch(r"[一-龥ァ-ヶぁ-んー]{2,4}", t) and t not in NOT_NAME:
                    names.add(t)
    names.add("ウエダ")
    return names


def mask_names(s, pattern):
    return pattern.sub("○○", s) if pattern else s


def supplement_sheets(sheetnames):
    """補足シート（No.348・QA-23_回答 など）を、(接頭辞, 番号) -> シート名のリストにする。"""
    found = {}
    for name in sheetnames:
        m = re.match(r"^(NO|QA)[.,\-]?\s*(\d+)(?:[-・](\d+))?", name.upper())
        if not m:
            continue
        kind = m[1]
        nums = {int(m[2])}
        if m[3] and name.upper().startswith(("NO.59", "QA-78")):  # 59-60 / 78・80 は2件分の共用シート
            nums.add(int(m[3]))
        for n in nums:
            found.setdefault((kind, n), []).append(name)
    return found


def clean_cell(s):
    return re.sub(r"\n{3,}", "\n\n", normalize(s)).strip()


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)
    supp = supplement_sheets(wb.sheetnames)

    data = {}
    for sheet in SHEETS:
        data[sheet] = [list(r[:19]) + [None] * (19 - len(r[:19]))
                       for r in wb[sheet].iter_rows(min_row=3, values_only=True) if r[0]]
    all_rows = [r for rows in data.values() for r in rows if text(r[C_STATUS]) == TARGET_STATUS]
    names = collect_names(all_rows)
    pattern = re.compile("|".join(sorted(map(re.escape, names), key=len, reverse=True)))

    for d in OUT.glob("*"):
        if d.is_dir():
            shutil.rmtree(d)
    summary = []
    for sheet, (prefix, supp_kind) in SHEETS.items():
        entries = []
        total = skipped = 0
        for r in data[sheet]:
            total += 1
            if text(r[C_STATUS]) != TARGET_STATUS:
                skipped += 1
                continue
            raw_no = text(r[C_NO])
            num = int(re.search(r"\d+", raw_no)[0])
            no = raw_no if raw_no.startswith(prefix) else f"{prefix}{raw_no}"
            fname = no.translate(BAD_FILENAME) + ".md"
            kind = unicodedata.normalize("NFKC", text(r[C_KIND]))
            func = " ／ ".join(x.strip() for x in text(r[C_FUNC]).splitlines() if x.strip())
            if func in ("-", "ー", "－"):
                func = ""
            title = " ".join(text(r[C_TITLE]).split())
            subs = supp.get((supp_kind, num), [])
            head = ["---",
                    f"出典: {XLSM.name} の「{sheet}」シート {no}",
                    f"シート: {sheet}", f"番号: {no}"]
            for k, v in (("区分", kind), ("機能名", func), ("件名", title),
                         ("発生部門", " ".join(text(r[C_DEPT]).split())),
                         ("発生日", text(r[C_DATE])), ("完了日", text(r[C_DONE])), ("状況", text(r[C_STATUS]))):
                if v:
                    head.append(f"{k}: {mask_names(normalize(v), pattern)}")
            if subs:
                head.append(f"補足シート: {'、'.join(subs)}（同Excel内。画面キャプチャ等の添付）")
            head += ["---", f"# {sheet}{no} {mask_names(normalize(title), pattern)}", ""]
            body = []
            for label, col in (("内容", C_BODY), ("対処・対策・回答案", C_ANSWER), ("結果（クローズ時）", C_RESULT)):
                if is_blank(r[col]):
                    continue
                body += [f"## {label}", "", mask_names(clean_cell(text(r[col])), pattern), ""]
            path = OUT / sheet / fname
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("\n".join(head + body), encoding="utf-8", newline="\n")
            entries.append(dict(file=fname, func=func, title=mask_names(normalize(title), pattern),
                                done=text(r[C_DONE]), num=num, supp=bool(subs)))
        # 機能名ごとに並べる索引
        entries.sort(key=lambda e: (e["func"], e["num"]))
        lines = [f"# {sheet}シート（状況=完了のみ）", "",
                 f"全{total}件のうち回答済み（状況={TARGET_STATUS}）の{len(entries)}件。機能名順。", "",
                 "| ファイル | 機能名 | 件名 | 完了日 | 補足 |", "|---|---|---|---|---|"]
        for e in entries:
            lines.append(f"| {e['file']} | {(e['func'] or '(未設定)').replace('|', '／')} | {e['title'].replace('|', '／')} "
                         f"| {e['done']} | {'有' if e['supp'] else ''} |")
        (OUT / sheet / "_INDEX.md").write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
        summary.append((sheet, total, len(entries), skipped))

    top = ["# 課題管理表（福原つながるプロジェクト）索引", "",
           f"元Excel: {XLSM.name}。実際の質疑応答・課題対応の記録。状況が「{TARGET_STATUS}」（回答済み）の行のみ。",
           "1件1ファイル。ファイル名＝番号（問合せ=`QA-n`／課題=`No.n`）で、元Excelのシート・番号をそのまま特定できる。",
           "発信者・対処者・確認者は伏せてある（本文中の氏名は「○○」）。",
           "「補足シート」が付く件は、元Excelに同じ番号のシート（画面キャプチャ等）がある。",
           "機能名・件名の一覧は各シートの `_INDEX.md`。", ""]
    for sheet, total, n, skipped in summary:
        top.append(f"- {sheet}シート: {n}件（全{total}件、未完了{skipped}件は除外） → `{sheet}/_INDEX.md`")
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "INDEX.md").write_text("\n".join(top) + "\n", encoding="utf-8", newline="\n")
    for sheet, total, n, skipped in summary:
        print(f"{sheet}: 全{total}件 / 出力{n}件 / 除外{skipped}件")
    print(f"マスク対象の氏名 {len(names)}語: {'、'.join(sorted(names))}")


if __name__ == "__main__":
    main()
